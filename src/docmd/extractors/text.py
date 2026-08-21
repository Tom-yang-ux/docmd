"""文本型提取器：PDF / Word(docx) / Excel(xlsx) / PPT(pptx) 原生文本。

对可复制文本优先文本提取，不调用视觉模型（成本控制）。
"""
from __future__ import annotations

import re
from typing import Optional

from ..core.document import ContentBlock
from ..core.models import BoundingBox
from ..core.normalize import normalize
from .base import BaseExtractor


def _block_id(prefix: str, n: int) -> str:
    return f"{prefix}{n:04d}"


class PdfTextExtractor(BaseExtractor):
    """提取 PDF 可复制文本；扫描页无文本时返回空（调用方转视觉）。"""
    engine = "text_extractor"

    def can_handle(self, file_type: str) -> bool:
        return file_type in ("pdf", ".pdf")

    def build_blocks(self, file_path: str) -> list[ContentBlock]:
        import pymupdf  # type: ignore
        blocks: list[ContentBlock] = []
        doc = pymupdf.open(file_path)
        n = 0
        for page_no, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if not text.strip():
                continue
            for line in text.splitlines():
                line = normalize(line.strip())
                if not line:
                    continue
                n += 1
                btype = self._classify(line)
                blocks.append(ContentBlock(
                    id=_block_id("b", n),
                    type=btype,
                    text=line,
                    page=page_no,
                    bbox=BoundingBox(page=page_no),
                ))
        doc.close()
        return blocks

    @staticmethod
    def _classify(line: str):
        from ..core.constants import BlockType
        if line.startswith("#") or (len(line) < 40 and re.match(r"^第[一二三四五六七八九十百]+[章节篇条]", line)):
            return BlockType.HEADING
        return BlockType.PARAGRAPH

    def is_scanned(self, file_path: str) -> bool:
        """判断是否为无文本层的扫描 PDF。"""
        import pymupdf  # type: ignore
        doc = pymupdf.open(file_path)
        has_text = False
        for page in doc:
            if page.get_text("text").strip():
                has_text = True
                break
        doc.close()
        return not has_text


class DocxExtractor(BaseExtractor):
    engine = "text_extractor"

    def can_handle(self, file_type: str) -> bool:
        return file_type in ("docx", ".docx")

    def build_blocks(self, file_path: str) -> list[ContentBlock]:
        from docx import Document as DocxDoc
        from ..core.constants import BlockType
        d = DocxDoc(file_path)
        blocks: list[ContentBlock] = []
        n = 0
        for para in d.paragraphs:
            t = normalize(para.text).strip()
            if not t:
                continue
            n += 1
            btype = BlockType.HEADING if para.style and para.style.name.startswith("Heading") else BlockType.PARAGRAPH
            blocks.append(ContentBlock(id=_block_id("b", n), type=btype, text=t))
        # 表格
        for tbl in d.tables:
            n += 1
            rows = []
            header = []
            for r_i, row in enumerate(tbl.rows):
                vals = [normalize(c.text).strip() for c in row.cells]
                if r_i == 0:
                    header = vals
                else:
                    rows.append(vals)
            blocks.append(ContentBlock(
                id=_block_id("b", n), type=BlockType.TABLE,
                text=f"表格 {n}", meta={"header": header, "rows": rows},
            ))
        return blocks


class XlsxExtractor(BaseExtractor):
    engine = "text_extractor"

    def can_handle(self, file_type: str) -> bool:
        return file_type in ("xlsx", ".xlsx")

    def build_blocks(self, file_path: str) -> list[ContentBlock]:
        import openpyxl
        from ..core.constants import BlockType
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        blocks: list[ContentBlock] = []
        n = 0
        for ws in wb.worksheets:
            n += 1
            header = []
            rows = []
            for r_i, row in enumerate(ws.iter_rows(values_only=True)):
                vals = ["" if v is None else normalize(str(v)).strip() for v in row]
                if r_i == 0:
                    header = vals
                else:
                    if any(vals):
                        rows.append(vals)
            blocks.append(ContentBlock(
                id=_block_id("b", n), type=BlockType.TABLE,
                text=f"工作表: {ws.title}",
                meta={"header": header, "rows": rows},
            ))
        wb.close()
        return blocks


class PptxExtractor(BaseExtractor):
    engine = "text_extractor"

    def can_handle(self, file_type: str) -> bool:
        return file_type in ("pptx", ".pptx")

    def build_blocks(self, file_path: str) -> list[ContentBlock]:
        from pptx import Presentation  # type: ignore
        prs = Presentation(file_path)
        blocks: list[ContentBlock] = []
        n = 0
        for slide_no, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    t = normalize(para.text).strip()
                    if not t:
                        continue
                    n += 1
                    blocks.append(ContentBlock(
                        id=_block_id("b", n), type=BlockType.PARAGRAPH, text=t, page=slide_no))
        return blocks


EXTRACTORS: list[BaseExtractor] = [
    PdfTextExtractor(),
    DocxExtractor(),
    XlsxExtractor(),
    PptxExtractor(),
]


def get_text_extractor(file_type: str) -> Optional[BaseExtractor]:
    for ext in EXTRACTORS:
        if ext.can_handle(file_type):
            return ext
    return None
